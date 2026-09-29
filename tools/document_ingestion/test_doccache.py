"""Tiny synthetic inputs only; no book or third-party document ingestion."""
import contextlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import doccache as dc


def tiny_pdf(path):
    # One page with a single text operand; stdlib fixture, not a publication.
    content = b"BT /F1 12 Tf 72 720 Td (Synthetic PDF Needle) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>",
               b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
               b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{x:010d} 00000 n \n".encode() for x in offsets[1:])
    data += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(data)


def tiny_docx(path):
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        z.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Synthetic DOCX Needle</w:t></w:r></w:p></w:body></w:document>')


class CacheTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base / "cache"
        self.source = self.base / "fixture.md"
        self.source.write_text("# Intro\n\nOriginal\n\n## Runtime\n\nGOMAXPROCS needle\n", encoding="utf-8")

    def test_miss_hit_invalidation_and_search(self):
        with patch.object(dc, "convert", wraps=dc.convert) as conversion:
            status, doc = dc.ingest(self.source, self.root)
            self.assertEqual(status, "CACHE_MISS")
            folder, index = dc.read_index(self.root, doc)
            self.assertTrue((folder / "source.json").exists())
            self.assertTrue((folder / "document.md").exists())
            self.assertTrue(index["chunks"])
            before = (folder / "source.json").read_bytes()
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_HIT", doc))
            self.assertEqual(conversion.call_count, 1)
            self.assertEqual(before, (folder / "source.json").read_bytes())
            for mode, query in [("text", "gomaxprocs"), ("exact", "GOMAXPROCS"), ("heading", "Runtime"), ("regex", "GOMAX.*needle")]:
                hits = dc.search(self.root, doc, query, mode)
                self.assertEqual(len(hits), 1)
                self.assertEqual(hits[0]["heading_path"], ["Intro", "Runtime"])
            self.assertFalse(dc.search(self.root, doc, "gomaxprocs", "exact"))
            timestamp = self.source.stat().st_mtime_ns
            self.source.write_text("# Changed\n\nNEW needle\n", encoding="utf-8")
            os.utime(self.source, ns=(timestamp, timestamp))
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_INVALID", doc))
            self.assertEqual(conversion.call_count, 2)
            after = json.loads((folder / "source.json").read_text())
            self.assertNotEqual(json.loads(before)["sha256"], after["sha256"])
            self.assertFalse(dc.search(self.root, doc, "GOMAXPROCS"))

    def test_converter_version_invalidation(self):
        with patch.object(dc, "version", return_value="1.0"):
            status, doc = dc.ingest(self.source, self.root)
            self.assertEqual(status, "CACHE_MISS")
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_HIT", doc))
        with patch.object(dc, "version", return_value="2.0"):
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_INVALID", doc))
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_HIT", doc))

    def test_converter_name_invalidation(self):
        _, doc = dc.ingest(self.source, self.root)
        with patch.object(dc, "CONVERTER", "alternate-converter"), patch.object(dc, "version", return_value="1"):
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_INVALID", doc))
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_HIT", doc))

    def test_pipeline_and_schema_invalidation(self):
        _, doc = dc.ingest(self.source, self.root)
        with patch.object(dc, "PIPELINE_VERSION", dc.PIPELINE_VERSION + 1):
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_INVALID", doc))
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_HIT", doc))
        with patch.object(dc, "SCHEMA", dc.SCHEMA + 1):
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_INVALID", doc))
            self.assertEqual(dc.ingest(self.source, self.root), ("CACHE_HIT", doc))

    def test_markdown_contains_only_normalized_source(self):
        raw = "# Source\r\n\r\nOnly extracted prose.\r\n"
        self.source.write_bytes(raw.encode("utf-8"))
        _, doc = dc.ingest(self.source, self.root)
        folder, index = dc.read_index(self.root, doc)
        markdown = (folder / "document.md").read_text(encoding="utf-8")
        self.assertEqual(markdown, dc.normalize(raw))
        self.assertNotIn("UNKNOWN/UNEXTRACTED", markdown)
        self.assertTrue(index["warnings"])
        source = json.loads((folder / "source.json").read_text(encoding="utf-8"))
        self.assertEqual(source["sha256"], dc.file_digest(self.source))

    def test_chunk_boundaries_and_offsets(self):
        code = "```go\n" + "println(1)\n" * 100 + "```\n"
        table = "| A | B |\n| --- | --- |\n" + "| value | other |\n" * 100
        text = "# Title\n\n" + "paragraph\n\n" * 50 + code + "\n" + table + "\n## Next\n\nend\n"
        chunks, headings = dc.chunk_document(text, target=200)
        slices = [text[c["char_start"]:c["char_end"]] for c in chunks]
        self.assertEqual("".join(slices), text)
        self.assertEqual(sum(code in s for s in slices), 1)
        self.assertEqual(sum(table in s for s in slices), 1)
        self.assertEqual(headings[1]["parent"], 0)
        for c, content in zip(chunks, slices):
            self.assertEqual(c["sha256"], dc.digest(content.encode()))
            self.assertIsNone(c["source_page_start"])

    def test_errors_never_publish_false_success(self):
        for name, data in [("broken.pdf", b"not PDF"), ("broken.docx", b"not ZIP"), ("broken.pptx", b"not ZIP"), ("unsupported.exe", b"x"), ("empty.txt", b"")]:
            source = self.base / name
            source.write_bytes(data)
            with self.subTest(name=name), contextlib.redirect_stderr(io.StringIO()):
                self.assertEqual(dc.main(["--cache-root", str(self.root), "ingest", str(source)]), 1)
                self.assertFalse((self.root / dc.document_id(source.resolve())).exists())
        self.assertFalse(list(self.root.glob("*.lock")))

    def test_failed_reconversion_preserves_cache(self):
        _, doc = dc.ingest(self.source, self.root)
        previous = (self.root / doc / "document.md").read_bytes()
        self.source.write_text("changed", encoding="utf-8")
        with self.assertRaises(RuntimeError):
            dc.ingest(self.source, self.root, converter=lambda p: (_ for _ in ()).throw(RuntimeError("failure")))
        self.assertEqual(previous, (self.root / doc / "document.md").read_bytes())

    def test_unicode_and_indented_code(self):
        code = "    print('x')\n\n    print('y')\n"
        text = "# Tiếng Việt\n\n" + code + "\nVăn bản có dấu.\n"
        self.source.write_text(text, encoding="utf-8")
        _, doc = dc.ingest(self.source, self.root)
        self.assertTrue(dc.search(self.root, doc, "TIẾNG VIỆT", "heading"))
        chunks, _ = dc.chunk_document(text, target=10)
        self.assertEqual(sum(code in text[c["char_start"]:c["char_end"]] for c in chunks), 1)

    def test_markdown_structure_survives_conversion(self):
        text = "# Heading\n\n- one\n- two\n\n```go\n// # not a heading\nprintln(1)\n```\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n"
        self.source.write_text(text, encoding="utf-8")
        _, doc = dc.ingest(self.source, self.root)
        folder, index = dc.read_index(self.root, doc)
        self.assertTrue((folder / "document.md").read_text(encoding="utf-8").startswith(text))
        self.assertEqual(len(index["heading_tree"]), 1)

    def test_source_mutation_and_writer_lock(self):
        def mutate(path):
            path.write_text("Changed mid-conversion", encoding="utf-8")
            return "stale output"
        with self.assertRaisesRegex(ValueError, "Source changed"):
            dc.ingest(self.source, self.root, converter=mutate)
        doc = dc.document_id(self.source.resolve())
        self.assertFalse((self.root / doc).exists())
        lock = self.root / (doc + ".lock")
        lock.touch()
        with self.assertRaises(FileExistsError):
            dc.ingest(self.source, self.root)
        self.assertTrue(lock.exists())

    def test_corrupt_chunk_regenerates(self):
        _, doc = dc.ingest(self.source, self.root)
        (self.root / doc / "chunks/0001.md").write_text("corrupted")
        self.assertEqual(dc.ingest(self.source, self.root)[0], "CACHE_INVALID")

    def test_corrupt_document_and_missing_or_extra_chunks_regenerate(self):
        _, doc = dc.ingest(self.source, self.root)
        folder = self.root / doc
        (folder / "document.md").write_text("tampered")
        self.assertEqual(dc.ingest(self.source, self.root)[0], "CACHE_INVALID")
        (folder / "chunks/extra.md").write_text("extra")
        self.assertEqual(dc.ingest(self.source, self.root)[0], "CACHE_INVALID")
        (folder / "chunks/0001.md").unlink()
        self.assertEqual(dc.ingest(self.source, self.root)[0], "CACHE_INVALID")

    def test_cache_symlinks_never_validate_or_escape(self):
        _, doc = dc.ingest(self.source, self.root)
        folder = self.root / doc
        chunk = folder / "chunks/0001.md"
        chunk.unlink()
        outside = self.base / "outside.md"
        outside.write_text("do not read", encoding="utf-8")
        chunk.symlink_to(outside)
        self.assertEqual(dc.ingest(self.source, self.root)[0], "CACHE_INVALID")

    def test_search_rejects_tampered_index_and_symlink(self):
        _, doc = dc.ingest(self.source, self.root)
        folder = self.root / doc
        index_path = folder / "index.json"
        original = index_path.read_text(encoding="utf-8")
        data = json.loads(original)
        data["chunks"][0]["id"] = "../../outside"
        index_path.write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "integrity"):
            dc.search(self.root, doc, "Original")
        index_path.write_text(original, encoding="utf-8")
        chunk = folder / "chunks/0001.md"
        chunk.unlink()
        outside = self.base / "outside.md"
        outside.write_text("secret", encoding="utf-8")
        chunk.symlink_to(outside)
        with self.assertRaisesRegex(ValueError, "integrity"):
            dc.search(self.root, doc, "secret")

    def test_real_pdf_docx_pptx_conversion(self):
        from pptx import Presentation
        pdf = self.base / "tiny.pdf"
        tiny_pdf(pdf)
        docx = self.base / "tiny.docx"
        tiny_docx(docx)
        pptx = self.base / "tiny.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Synthetic PPTX Needle"
        presentation.save(pptx)
        for source in (pdf, docx, pptx):
            with self.subTest(format=source.suffix):
                status, doc = dc.ingest(source, self.root)
                self.assertEqual(status, "CACHE_MISS")
                self.assertTrue(dc.search(self.root, doc, "Needle"))
                self.assertEqual(dc.ingest(source, self.root)[0], "CACHE_HIT")

    def test_cli_show_and_invalid_regex(self):
        _, doc = dc.ingest(self.source, self.root)
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            self.assertEqual(dc.main(["--cache-root", str(self.root), "show", doc, "0001"]), 0)
        self.assertIn("Original", output.getvalue())
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(dc.main(["--cache-root", str(self.root), "search", doc, "[", "--mode", "regex"]), 1)
            self.assertEqual(dc.main(["--cache-root", str(self.root), "show", "../escape", "0001"]), 1)


if __name__ == "__main__":
    unittest.main()
